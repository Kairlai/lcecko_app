import base64
from datetime import datetime, timedelta
import io
from io import BytesIO
import os
import random
import urllib.parse
import pandas as pd
import qrcode
import requests
import json
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

FILE_PATH = "objednavky_lcecko.csv"
SETTINGS_PATH = "nastaveni_lcecko.json"
LOGO_PATH = "logo.jpg"

# --- BANKOVNÍ ÚDAJE L-CÉČKA ---
BANK_ACCOUNT = "4610796083"
BANK_CODE = "0800"
BANK_IBAN = "CZ2208000000004610796083"

GITHUB_TOKEN = st.secrets.get("GITHUB_TOKEN", "")
GITHUB_REPO = st.secrets.get("GITHUB_REPO", "")
ADMIN_PASSWORD = st.secrets.get("ADMIN_PASSWORD", "heslo1234")

# --- NASTAVENÍ E-MAILU ZE SECRETS ---
EMAIL_SENDER = st.secrets.get("EMAIL_SENDER", "")
EMAIL_PASSWORD = st.secrets.get("EMAIL_PASSWORD", "")
EMAIL_RECEIVER = st.secrets.get("EMAIL_RECEIVER", "")
SMTP_SERVER = st.secrets.get("SMTP_SERVER", "smtp.centrum.cz")
SMTP_PORT = int(st.secrets.get("SMTP_PORT", 465))

st.set_page_config(page_title="L-Céčko | Objednávkový systém", layout="wide", page_icon="🥗")

# Vizuální styl L-Céčka
st.markdown("""
    <style>
    #MainMenu {visibility: hidden;}
    [data-testid="stHeaderActionElements"] {visibility: hidden;}
    header {background-color: transparent !important;}
    footer {visibility: hidden;}
    
    h1, h2, h3, .main-header {
        color: #4A7833 !important;
    }

    div[data-testid="stColumn"] {
        background: #ffffff;
        padding: 20px;
        border-radius: 12px;
        box-shadow: 0 4px 12px rgba(0,0,0,0.08);
        border-top: 5px solid #4A7833;
    }
    
    div.stButton > button:first-child {
        background-color: #4A7833;
        color: white;
        border: none;
        border-radius: 8px;
        font-weight: bold;
    }
    div.stButton > button:first-child:hover {
        background-color: #385a26;
        color: white;
    }
    
    .summary-card {
        background-color: #e6f0e1;
        padding: 15px;
        border-radius: 8px;
        border-left: 5px solid #4A7833;
        margin-bottom: 20px;
    }
    
    .filter-box {
        background-color: #f8f9fa;
        padding: 15px 15px 5px 15px;
        border-radius: 8px;
        border: 1px solid #dee2e6;
        margin-bottom: 20px;
    }
    
    /* STYLY PRO BĚŽÍCÍ TITULKY */
    .marquee-container {
        width: 100%;
        overflow: hidden;
        background-color: #ffeb3b;
        color: #333;
        padding: 8px 0;
        font-weight: bold;
        border-radius: 5px;
        margin-bottom: 15px;
        white-space: nowrap;
        box-shadow: 0 2px 5px rgba(0,0,0,0.1);
    }
    .marquee-content {
        display: inline-block;
        padding-left: 100%;
        animation: marquee 20s linear infinite;
    }
    @keyframes marquee {
        0%   { transform: translate(0, 0); }
        100% { transform: translate(-100%, 0); }
    }
    </style>
""", unsafe_allow_html=True)

def odeslat_email_upozorneni(id_obj, jmeno, prijmeni, adresa, program, delka, datum_od, cena, poznamka):
    if not EMAIL_SENDER or not EMAIL_PASSWORD or not EMAIL_RECEIVER:
        print("E-mailové údaje chybí v konfiguraci.")
        return False
    try:
        msg = MIMEMultipart()
        msg['From'] = EMAIL_SENDER
        msg['To'] = EMAIL_RECEIVER
        msg['Subject'] = f"🥗 Nová objednávka #{id_obj} - {jmeno} {prijmeni}"
        
        body = f"""Dobrý den,

přes objednávkový formulář L-Céčka právě dorazila nová objednávka!

Zákazník: {jmeno} {prijmeni}
Adresa pro kurýra: {adresa}
Poznámka pro kurýra / alergie: {poznamka if poznamka else '-'}

Menu / Program: {program}
Varianta / Délka: {delka}
Od kdy doručovat: {datum_od}
Celková cena k úhradě: {cena:,.0f} Kč

Objednávku si můžete detailně zkontrolovat v dispečinku aplikace.

Hezký den,
Váš systém L-Céčko
"""
        msg.attach(MIMEText(body, 'plain', 'utf-8'))
        
        if SMTP_PORT == 465:
            server = smtplib.SMTP_SSL(SMTP_SERVER, SMTP_PORT, timeout=5)
        else:
            server = smtplib.SMTP(SMTP_SERVER, SMTP_PORT, timeout=5)
            server.starttls()
            
        server.login(EMAIL_SENDER, EMAIL_PASSWORD)
        server.send_message(msg)
        server.quit()
        return True
    except Exception as e:
        print(f"Chyba pri odesilani e-mailu: {e}")
        return False

def get_headers():
    return {
        "Authorization": f"token {GITHUB_TOKEN}",
        "Accept": "application/vnd.github.v3+json"
    }

def nacti_nastaveni():
    vychozi = {
        "snidane": True, 
        "cukrovi": True, 
        "zakazkova": True, 
        "min_ks_catering": 1,
        "bezici_text": "🎄 Přijímáme objednávky na Vánoční cukroví do 31.10.2026! 🎄 Dotazy / Mimořádné přání: 775 045 554 nebo formulář na www.l-cecko.cz"
    }
    if not GITHUB_TOKEN or not GITHUB_REPO:
        if not os.path.exists(SETTINGS_PATH):
            with open(SETTINGS_PATH, "w") as f:
                json.dump(vychozi, f)
        with open(SETTINGS_PATH, "r") as f:
            return json.load(f), None

    try:
        url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{SETTINGS_PATH}"
        res = requests.get(url, headers=get_headers(), timeout=5)
        if res.status_code == 200:
            data = res.json()
            sha = data["sha"]
            content_str = base64.b64decode(data["content"]).decode("utf-8")
            loaded = json.loads(content_str)
            for k, v in vychozi.items():
                if k not in loaded:
                    loaded[k] = v
            return loaded, sha
    except Exception as e:
        print(f"Chyba při načítání nastavení z GitHubu: {e}")
    
    return vychozi, None

def uloz_nastaveni(nastaveni_dict, sha=None):
    obsah = json.dumps(nastaveni_dict)
    if not GITHUB_TOKEN or not GITHUB_REPO:
        with open(SETTINGS_PATH, "w") as f:
            f.write(obsah)
        return True

    try:
        url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{SETTINGS_PATH}"
        content_b64 = base64.b64encode(obsah.encode("utf-8")).decode("utf-8")
        payload = {"message": "Aktualizace nastavení", "content": content_b64}
        if sha:
            payload["sha"] = sha
        res = requests.put(url, headers=get_headers(), json=payload, timeout=5)
        return res.status_code in [200, 201]
    except Exception as e:
        print(f"Chyba při ukládání nastavení na GitHub: {e}")
        return False

def nacti_objednavky():
    default_columns = [
        "ID", "Datum_Vytvoreni", "Jmeno", "Prijmeni", "Telefon", "Email", 
        "Adresa", "Poznamka_Kuryr", "Kategorie", "Program", "Delka", 
        "Datum_Od", "Cena_Celkem", "Stav_Platby"
    ]
    
    if not GITHUB_TOKEN or not GITHUB_REPO:
        if not os.path.exists(FILE_PATH):
            df_empty = pd.DataFrame(columns=default_columns)
            df_empty.to_csv(FILE_PATH, index=False)
            return df_empty, None
        return pd.read_csv(FILE_PATH), None

    try:
        url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{FILE_PATH}"
        res = requests.get(url, headers=get_headers(), timeout=5)
        if res.status_code == 200:
            data = res.json()
            sha = data["sha"]
            content_str = base64.b64decode(data["content"]).decode("utf-8")
            return pd.read_csv(io.StringIO(content_str)), sha
    except Exception as e:
        print(f"Chyba při načítání objednávek z GitHubu: {e}")
        
    df_empty = pd.DataFrame(columns=default_columns)
    return df_empty, None

def uloz_objednavky(df, sha=None):
    if not GITHUB_TOKEN or not GITHUB_REPO:
        df.to_csv(FILE_PATH, index=False)
        return True

    try:
        url = f"https://api.github.com/repos/{GITHUB_REPO}/contents/{FILE_PATH}"
        csv_buffer = io.StringIO()
        df.to_csv(csv_buffer, index=False)
        content_b64 = base64.b64encode(csv_buffer.getvalue().encode("utf-8")).decode("utf-8")

        payload = {"message": "Aktualizace objednavek L-Cecko", "content": content_b64}
        if sha:
            payload["sha"] = sha
        res = requests.put(url, headers=get_headers(), json=payload, timeout=5)
        return res.status_code in [200, 201]
    except Exception as e:
        print(f"Chyba při ukládání objednávek na GitHub: {e}")
        return False

def vytvor_profi_excel(df, titulek="Objednávky"):
    wb = Workbook()
    ws = wb.active
    ws.title = titulek
    
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
        
    green_fill = PatternFill(start_color="4A7833", end_color="4A7833", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )
    
    for cell in ws[1]:
        cell.fill = green_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[col_letter].width = (max_length + 2)
        
    ws.auto_filter.ref = ws.dimensions
    
    output = BytesIO()
    wb.save(output)
    return output.getvalue()

def detekuj_alergie(text):
    """Skenuje poznámky a hledá klíčová slova k alergiím."""
    if pd.isna(text):
        return ""
    text_low = str(text).lower()
    klicova_slova = ["alerg", "lepek", "bezlep", "laktóz", "laktoz", "ořech", "mlék", "mlek", "česnek", "cibul", "ryb"]
    for slovo in klicova_slova:
        if slovo in text_low:
            return "🚨 ALERGIE/VÝJIMKA!"
    return ""

def vygeneruj_html_uctenku(id_obj, jmeno, prijmeni, adresa, telefon, program, delka, datum, cena):
    """Generátor krásné účtenky pro zákazníka."""
    html_obsah = f"""
    <!DOCTYPE html>
    <html lang="cs">
    <head>
    <meta charset="UTF-8">
    <title>Potvrzení objednávky L-Céčko</title>
    <style>
        body {{ font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif; color: #333; max-width: 600px; margin: 0 auto; padding: 40px 20px; background-color: #f9f9f9; }}
        .uctenka {{ background: #fff; padding: 30px; border-radius: 10px; box-shadow: 0 4px 15px rgba(0,0,0,0.05); border-top: 6px solid #4A7833; }}
        .header {{ text-align: center; border-bottom: 2px dashed #ddd; padding-bottom: 20px; margin-bottom: 20px; }}
        .header h1 {{ color: #4A7833; margin: 0 0 5px 0; font-size: 28px; }}
        .header p {{ color: #777; margin: 0; font-size: 14px; }}
        .section-title {{ font-size: 14px; color: #888; text-transform: uppercase; letter-spacing: 1px; margin-bottom: 10px; border-bottom: 1px solid #eee; padding-bottom: 5px; }}
        .info-row {{ display: flex; justify-content: space-between; margin-bottom: 8px; font-size: 15px; }}
        .info-row strong {{ font-weight: 600; }}
        .total-box {{ background-color: #e6f0e1; padding: 15px; border-radius: 8px; margin-top: 25px; text-align: center; }}
        .total-box h2 {{ margin: 0; color: #4A7833; font-size: 24px; }}
        .payment-info {{ margin-top: 25px; font-size: 14px; background: #fff8e1; padding: 15px; border-radius: 8px; border-left: 4px solid #ffc107; }}
        .footer {{ text-align: center; margin-top: 30px; font-size: 13px; color: #999; }}
    </style>
    </head>
    <body>
        <div class="uctenka">
            <div class="header">
                <h1>🥗 L-Céčko Plzeň</h1>
                <p>Potvrzení přijetí objednávky #{id_obj}</p>
            </div>
            
            <div class="section-title">Údaje zákazníka</div>
            <div class="info-row"><span>Zákazník:</span> <strong>{jmeno} {prijmeni}</strong></div>
            <div class="info-row"><span>Adresa:</span> <strong>{adresa}</strong></div>
            <div class="info-row"><span>Telefon:</span> <strong>{telefon}</strong></div>
            <br>
            
            <div class="section-title">Souhrn objednávky</div>
            <div class="info-row"><span>Program:</span> <strong>{program}</strong></div>
            <div class="info-row"><span>Varianta:</span> <strong>{delka}</strong></div>
            <div class="info-row"><span>Start rozvozu:</span> <strong>{datum}</strong></div>
            
            <div class="total-box">
                <span style="font-size: 14px; color: #666; text-transform: uppercase;">Celkem k úhradě</span>
                <h2>{cena} Kč</h2>
            </div>
            
            <div class="
